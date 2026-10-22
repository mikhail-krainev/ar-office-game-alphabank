"""Builds docs/presentation/alfa-office.pptx.

Usage: python3 build_deck.py <mockup_dir>
Run mockups.py into the same directory first.
"""

import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "presentation" / "alfa-office.pptx"
SHOTS = ROOT / "docs" / "screenshots"

RED = RGBColor(0xEF, 0x31, 0x24)
INK = RGBColor(0x12, 0x12, 0x14)
NIGHT = RGBColor(0x0E, 0x0E, 0x10)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xF3, 0xF4, 0xF5)
MUTED = RGBColor(0x6D, 0x70, 0x74)
HAIR = RGBColor(0xD6, 0xD8, 0xDB)
NIGHT_MUTED = RGBColor(0x9A, 0x9C, 0xA1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Arial"
SLIDE_W, SLIDE_H = 13.333, 7.5
M = 0.667  # side margin
RIGHT = SLIDE_W - M
TITLE_Y = 0.6
BODY_Y = 2.15  # content top under a two-line title
BODY_Y_SHORT = 1.75  # content top under a one-line title
SHADOW_PAD = 90  # px, must match mockups.py

SHORT_WORDS = "в во и к ко с со о об у на не за по из до от без для а но что как это при под над то же ли ни".split()


def nb(text: str) -> str:
    """Russian typesetting: glue short words to the next word, numbers to their units."""
    pattern = r"(?<![\w-])(" + "|".join(SHORT_WORDS) + r") "
    text = re.sub(pattern, lambda m: m.group(1) + "\u00a0", text, flags=re.IGNORECASE)
    text = re.sub(pattern, lambda m: m.group(1) + "\u00a0", text, flags=re.IGNORECASE)
    text = re.sub(r"(\d) (?=[а-яА-Я%×]|\d{3})", "\\1\u00a0", text)
    return text.replace(" — ", "\u00a0— ")


# ---------------------------------------------------------------- primitives


def background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, paragraphs, size=15, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.18, space_after=0, name=None):
    """paragraphs: str, or a list of str / list of (text, overrides) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, over in runs:
            r = p.add_run()
            r.text = nb(run_text)
            f = r.font
            f.name = FONT
            f.size = Pt(over.get("size", size))
            f.bold = over.get("bold", bold)
            f.color.rgb = over.get("color", color)
    return box


def hairline(slide, x1, y1, x2, y2, color=HAIR, weight=0.75):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(weight)
    no_effects(line)
    return line


def no_effects(shape):
    """Drop the theme style reference: it brings a shadow that PowerPoint and LibreOffice draw differently.
    Fill and line are always set explicitly."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def box(slide, x, y, w, h, fill, radius=None, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = radius / min(w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    no_effects(shape)
    return shape


def title(slide, words, number, dark=False):
    color = WHITE if dark else INK
    text(slide, M, TITLE_Y, 10.4, 1.0, words, size=34, bold=True, color=color, spacing=1.0, name="Title")
    text(slide, RIGHT - 1.0, TITLE_Y + 0.08, 1.0, 0.3, f"{number:02d}", size=11,
         color=NIGHT_MUTED if dark else MUTED, align=PP_ALIGN.RIGHT, name="Number")


class Mockups:
    scratch = "."

    def __init__(self, folder: Path):
        self.folder = folder
        Mockups.scratch = str(folder)

    def phone(self, slide, shot, x, y, h):
        """Places a phone so that its body (without shadow) has its top-left at x, y and height h.
        Parts outside the slide are cropped away."""
        return self._place(slide, self.folder / f"phone-{shot}.png", x, y, h)

    def browser(self, slide, shot, x, y, w):
        path = self.folder / f"browser-{shot}.png"
        iw, ih = Image.open(path).size
        h = w * (ih - 2 * SHADOW_PAD) / (iw - 2 * SHADOW_PAD)
        return self._place(slide, path, x, y, h)

    @staticmethod
    def body_width(h):
        return h * 838 / 1756

    def _place(self, slide, path, x, y, h):
        iw, ih = Image.open(path).size
        scale = h / (ih - 2 * SHADOW_PAD)
        pad = SHADOW_PAD * scale
        left, top = x - pad, y - pad
        width, height = iw * scale, ih * scale
        crop_r = max(0.0, (left + width - SLIDE_W) / width)
        crop_b = max(0.0, (top + height - SLIDE_H) / height)
        pic = slide.shapes.add_picture(
            str(path), Inches(left), Inches(top), Inches(width * (1 - crop_r)), Inches(height * (1 - crop_b))
        )
        pic.crop_right = crop_r
        pic.crop_bottom = crop_b
        pic.name = path.stem
        return pic


def caption(slide, x, y, w, head, body, dark=False):
    text(slide, x, y, w, 0.6, [[(head, {"bold": True, "size": 13, "color": WHITE if dark else INK})],
                                [(body, {"size": 11, "color": NIGHT_MUTED if dark else MUTED})]],
         spacing=1.12, space_after=2)


def rows(slide, x, y, w, label_w, items, row_h, size=15, label_size=15, top_rule=True):
    """Editorial list: bold label on the left, text on the right, hairlines between rows."""
    if top_rule:
        hairline(slide, x, y, x + w, y)
    for i, (label, body) in enumerate(items):
        ry = y + i * row_h
        text(slide, x, ry + 0.16, label_w - 0.2, row_h - 0.2, label, size=label_size, bold=True)
        text(slide, x + label_w, ry + 0.16, w - label_w, row_h - 0.2, body, size=size, color=INK)
        hairline(slide, x, ry + row_h, x + w, ry + row_h)


def prose(slide, x, y, w, items, size=15, color=INK, dark=False):
    """A short column of paragraphs with air between them."""
    return text(slide, x, y, w, 4.5, items, size=size, color=WHITE if dark else color, spacing=1.22,
                space_after=14)


# ---------------------------------------------------------------- slides


def slide_cover(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, RED)
    text(s, M, TITLE_Y, 5.6, 0.3, "Студенческий концепт-проект для Альфа-Банка", size=13, color=WHITE)
    text(s, M, 1.55, 5.2, 2.6, "Альфа\nОфис", size=96, bold=True, color=WHITE, spacing=0.92)
    text(s, M, 4.85, 4.6, 1.2, "Браузерная игра для телефона, ради которой хочется приехать в офис",
         size=22, color=WHITE, spacing=1.15)
    text(s, M, 6.62, 4.0, 0.3, "УрФУ, 2026", size=13, color=WHITE)
    h = 6.6
    w = mock.body_width(h)
    xb = RIGHT - w
    mock.phone(s, "06-office", xb - 0.35 - w, 1.35, h)
    mock.phone(s, "51-ar-selfie-crown", xb, 2.05, h)
    s.notes_slide.notes_text_frame.text = (
        "Альфа Офис — браузерная игра для телефона. Её цель простая: дать сотруднику причину приехать в офис. "
        "Все задания выполняются физически в офисе, а играть можно только из корпоративного Wi-Fi."
    )


def slide_idea(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, PAPER)
    title(s, "Игра, в которую нельзя\nиграть из дома", 2)
    text(s, M, BODY_Y, 7.7, 1.3,
         "Каждый рабочий день сотрудник получает свой набор заданий. Выполнить их можно только в офисе: "
         "нужны камера, шагомер, микрофон и QR-коды на дверях комнат. За задания начисляются альфа-коины.",
         size=17, color=INK, spacing=1.22)
    rows(s, M, 3.85, 8.1, 1.75, [
        ("Прийти", "Отметиться на ресепшне. QR-код на экране меняется каждые 30 секунд, скриншот не сработает."),
        ("Сыграть", "Пройти мини-игры в комнатах офиса: 10 приседаний, 120 шагов по лестнице, фото с коллегой."),
        ("Обменять", "Потратить коины: мерч, промокоды на такси и кофе, парковка у входа, выходной день."),
    ], row_h=0.95)
    h = 5.75
    mock.phone(s, "08-tasks", RIGHT - mock.body_width(h), 1.15, h)
    s.notes_slide.notes_text_frame.text = (
        "Схема из трёх шагов. Важно, что каждое действие подтверждает сервер: клиент сам ничего не начисляет."
    )


def slide_day(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, MIST)
    title(s, "Один игровой день повторяет рабочий", 3)
    shots = [
        ("01-morning", "Утро", "Будильник, серия дней и множитель наград"),
        ("02-home", "Квартира", "Гардероб, гараж с машинами, выход"),
        ("05-commute", "Дорога", "Катсцена зависит от выбранной машины"),
        ("06-office", "Офис", "Задания, коллеги, магазин наград"),
        ("12-evening", "Вечер", "«Поехать домой» закрывает визит"),
    ]
    gutter = 0.333
    w = (RIGHT - M - gutter * 4) / 5
    h = w * 1756 / 838
    y = BODY_Y_SHORT
    for i, (shot, head, body) in enumerate(shots):
        x = M + i * (w + gutter)
        mock.phone(s, shot, x, y, h)
        caption(s, x, y + h + 0.2, w, head, body)
    s.notes_slide.notes_text_frame.text = (
        "Утро и дорога — катсцены в стиле комикса. Квартира и офис — карта в виде 3/4 сверху. "
        "После отметки на ресепшне визит уже засчитан, дальше задания только добавляют коины."
    )


TASKS = [
    ("Отметка на ресепшне", "Сканирует QR-код на экране у входа", "камера", "10"),
    ("Лестница, не лифт", "Проходит 120 шагов по офису за пять минут", "шагомер", "25"),
    ("Найди предмет", "Показывает камере три офисные вещи из 25", "объекты", "25"),
    ("Утренняя зарядка", "Делает 10 приседаний перед камерой", "поза", "30 ×2"),
    ("Скороговорка", "Произносит две скороговорки вслух", "речь", "20 ×3"),
    ("Фото с коллегой", "Двое в кадре, коллега подтверждает у себя", "лица, AR", "30 ×2"),
    ("Охота на баги", "Ловит баги, поворачиваясь с телефоном", "гироскоп, AR", "20"),
    ("Настроение дня", "Рисует настроение и примеряет его на лоб", "лица, AR", "15"),
    ("Три правды, две лжи", "Коллега угадывает факты о нём", "QR", "35 ×2"),
]


def slide_tasks(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, PAPER)
    title(s, "Мини-игры работают на датчиках телефона", 4)
    x, w = M, 8.35
    cols = [(0, 2.35), (2.35, 3.95), (6.3, 1.25), (7.55, 0.8)]
    heads = ["Задание", "Что делает игрок", "Датчик", "Коины"]
    y, row_h = BODY_Y_SHORT, 0.44
    for (cx, cw), head in zip(cols, heads):
        align = PP_ALIGN.RIGHT if head == "Коины" else PP_ALIGN.LEFT
        text(s, x + cx, y, cw, 0.3, head, size=11, color=MUTED, align=align)
    y += 0.34
    hairline(s, x, y, x + w, y, color=INK, weight=1)
    for i, row in enumerate(TASKS):
        ry = y + i * row_h
        for (cx, cw), value, col in zip(cols, row, range(4)):
            align = PP_ALIGN.RIGHT if col == 3 else PP_ALIGN.LEFT
            text(s, x + cx, ry, cw - 0.15 if col < 3 else cw, row_h, value, size=13, bold=col == 0,
                 color=INK if col != 2 else MUTED, align=align, anchor=MSO_ANCHOR.MIDDLE)
        hairline(s, x, ry + row_h, x + w, ry + row_h)
    text(s, x, y + len(TASKS) * row_h + 0.22, w, 0.6,
         "×2 и ×3 — множитель сложности, такое задание можно пропустить. Если датчика нет или игрок не дал "
         "разрешение, запускается упрощённая экранная версия.", size=11, color=MUTED)
    h = 5.15
    mock.phone(s, "17-find-object", RIGHT - mock.body_width(h), BODY_Y_SHORT, h)
    s.notes_slide.notes_text_frame.text = (
        "Задания физические, поэтому их нельзя сделать на диване. Набор на день у каждого свой: "
        "два рабочих дня подряд делят все задания пополам."
    )


def slide_ar(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NIGHT)
    title(s, "Камера показывает офис,\nигра дорисовывает остальное", 5, dark=True)
    text(s, M, BODY_Y, 3.55, 2.6, [
        "Поверх живой камеры появляются неоновые объекты: очки и корона на лицах, скелет при приседаниях, "
        "баги вокруг игрока.",
        "Распознавание работает прямо в браузере на MediaPipe. Кадры не покидают телефон.",
    ], size=15, color=WHITE, spacing=1.22, space_after=12)
    shots = [
        ("47-bug-hunt", "Охота на баги", "Баги висят вокруг игрока"),
        ("51-ar-selfie-crown", "Фото с коллегой", "Улыбка даёт корону"),
        ("52-ar-squats", "Утренняя зарядка", "Скелет и счёт приседаний"),
    ]
    h, gutter = 4.2, 0.333
    w = mock.body_width(h)
    x0 = RIGHT - 3 * w - 2 * gutter
    for i, (shot, head, body) in enumerate(shots):
        x = x0 + i * (w + gutter)
        mock.phone(s, shot, x, BODY_Y, h)
        caption(s, x, BODY_Y + h + 0.2, w, head, body, dark=True)
    s.notes_slide.notes_text_frame.text = (
        "AR-слой рисуется в Godot поверх видео с камеры. Лица только считаются, никого не опознают, "
        "снимки не сохраняются."
    )


def slide_colleagues(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, PAPER)
    title(s, "По офису ходят коллеги,\nкоторые сегодня пришли", 6)
    shots = ["44-office-colleagues", "45-colleague-profile", "46-colleague-story"]
    h, gutter = 4.75, 0.333
    w = mock.body_width(h)
    for i, shot in enumerate(shots):
        mock.phone(s, shot, M + i * (w + gutter), BODY_Y, h)
    tx = M + 3 * w + 2 * gutter + 0.6
    tw = RIGHT - tx
    prose(s, tx, BODY_Y, tw, [
        ("Каждый персонаж одет в образ, который его владелец собрал в гардеробе. Розовое кольцо "
             "у аватарки — есть история «Настроение дня»."),
        ("Нажатие на коллегу открывает его карточку: департамент, статус, серия дней в офисе."),
        ("Напарника для фото, «Кофе вслепую» и «Бинго знакомств» сервер выбирает только среди тех, "
             "кто сегодня отметился в этом же офисе."),
    ])
    s.notes_slide.notes_text_frame.text = (
        "Социальная часть — главный аргумент приехать: встречи засчитываются только вживую, "
        "по QR-коду с телефона коллеги, который меняется каждые 10 секунд."
    )


def slide_presence(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, MIST)
    title(s, "Присутствие проверяет сервер,\nа не честное слово", 7)
    factors = [
        ("1", "Ротируемый QR у ресепшна",
         "Код подписан секретом, который есть только у сервера. Живёт 30 секунд и не принимается повторно. "
         "Между кодами «Вход» и «Выход» игрок считается в офисе."),
        ("2", "Wi-Fi офиса",
         "Сервер сверяет внешний IP запроса с сетями офисов. Из дома и мобильного интернета любой игровой "
         "запрос получает отказ."),
    ]
    y = BODY_Y
    for num, head, body in factors:
        text(s, M, y - 0.08, 0.6, 0.8, num, size=44, bold=True, color=RED, spacing=0.9)
        text(s, M + 0.75, y, 5.2, 0.4, head, size=19, bold=True)
        text(s, M + 0.75, y + 0.48, 5.2, 1.3, body, size=15, spacing=1.22)
        y += 2.0
    h = 4.3
    w = mock.body_width(h)
    px = RIGHT - w
    mock.phone(s, "14-qr-scanner", px, BODY_Y, h)
    office_screen(s, px - 0.5 - w, BODY_Y, w, h)
    caption(s, px - 0.5 - w, BODY_Y + h + 0.2, w, "Экран у ресепшна", "Коды входа и выхода")
    caption(s, px, BODY_Y + h + 0.2, w, "Телефон сотрудника", "Сканер QR в игре")
    s.notes_slide.notes_text_frame.text = (
        "Два фактора проверяются на сервере. Если сотрудник открыл игру в Wi-Fi офиса, но забыл код, "
        "серия сохраняется, а игра напоминает отсканировать QR."
    )


def office_screen(s, x, y, w, h):
    """The reception screen: a dark display with the rotating code, without the phone's cut-outs."""
    bezel = 0.1
    panel = box(s, x, y, w, h, RGBColor(0x05, 0x05, 0x06), radius=0.1, line=RGBColor(0x3A, 0x3A, 0x3E))
    panel.name = "reception-screen"
    shot = Image.open(SHOTS / "16-office-screen.png").convert("RGB")
    screen = shot.crop((0, 140, shot.width, shot.height - 110))
    path = Path(Mockups.scratch) / "reception-screen.png"
    screen.save(path)
    inner_w = w - 2 * bezel
    inner_h = min(h - 2 * bezel, inner_w * screen.height / screen.width)
    pic = s.shapes.add_picture(str(path), Inches(x + bezel), Inches(y + (h - inner_h) / 2),
                               Inches(inner_w), Inches(inner_h))
    pic.name = "reception-screen-content"


def slide_economy(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, PAPER)
    title(s, "Коины начисляет только сервер", 8)
    text(s, M, BODY_Y_SHORT, 6.0, 0.5, [[("награда = база × сложность × серия", {"bold": True})]], size=24)
    parts = [("база", "10–40 коинов"), ("сложность", "×1, ×2 или ×3"), ("серия", "от ×1,0 до ×2,0, +0,1 за день")]
    for i, (head, body) in enumerate(parts):
        text(s, M + i * 2.0, BODY_Y_SHORT + 0.65, 1.85, 0.7, [[(head, {"size": 11, "color": MUTED})], [(body, {"size": 13})]],
             spacing=1.15)
    rows(s, M, 3.55, 5.9, 1.25, [
        ("500", "приветственный бонус при первом входе"),
        ("700", "лимит начислений в день"),
        ("0", "коинов клиент может начислить сам: каждое изменение баланса проверяет сервер"),
    ], row_h=0.78, label_size=20)
    chart_x = 7.35
    text(s, chart_x, BODY_Y_SHORT, RIGHT - chart_x, 0.4, "Пропуски офиса стоят коинов", size=19, bold=True)
    text(s, chart_x, BODY_Y_SHORT + 0.45, RIGHT - chart_x, 0.7,
         "Процент от баланса за каждый рабочий день без офиса. Первый пропуск сжигает только серию; "
         "отпуск и больничный не считаются.", size=13, color=MUTED, spacing=1.18)
    penalty_chart(s, chart_x, 3.3, RIGHT - chart_x, 3.4)
    s.notes_slide.notes_text_frame.text = (
        "Экономика целиком на сервере, в кошельке Nakama. Все RPC, меняющие баланс, идемпотентны."
    )


def penalty_chart(s, x, y, w, h):
    data = CategoryChartData(number_format='0"%"')
    data.categories = ["1 день", "2 дня", "3–4", "5–9", "10+"]
    data.add_series("Штраф", (0, 3, 5, 10, 15))
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    frame.name = "penalty-chart"
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(12)
    chart.font.color.rgb = MUTED
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = '0"%"'
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(14)
    labels.font.bold = True
    labels.font.color.rgb = INK
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RED
    value_axis = chart.value_axis
    value_axis.visible = False
    value_axis.has_major_gridlines = False
    value_axis.maximum_scale = 17
    value_axis.minimum_scale = 0
    category_axis = chart.category_axis
    category_axis.format.line.color.rgb = HAIR
    category_axis.has_major_gridlines = False
    category_axis.tick_labels.font.size = Pt(12)
    category_axis.tick_labels.font.color.rgb = MUTED


def slide_shop(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, MIST)
    title(s, "Коины меняются на то,\nчто ценят в офисе", 9)
    columns = [
        ("10-shop", "Магазин",
         "Шесть случайных предметов в день, от стикеров до выходного. Дорогие награды подтверждает "
         "человек."),
        ("26-raffle", "Розыгрыш парковки",
         "Раз в месяц — место на тёплом паркинге. Билет стоит 1 500 коинов, нужен 8 дней в офисе из 30."),
        ("27-roulette", "Честный итог",
         "Победителя выбирает сервер: хеш зерна публикуется заранее, итог может проверить любой."),
    ]
    gutter = 0.333
    col_w = (RIGHT - M - 2 * gutter) / 3
    for i, (shot, head, body) in enumerate(columns):
        x = M + i * (col_w + gutter)
        text(s, x, BODY_Y, col_w - 0.3, 1.3, [[(head, {"bold": True, "size": 17})], [(body, {"size": 14})]],
             spacing=1.2, space_after=6)
        mock.phone(s, shot, x, 3.85, 6.0)
    s.notes_slide.notes_text_frame.text = (
        "Купленный мерч открывает такие же вещи в гардеробе. Рулетка в день итогов — только анимация, "
        "результат уже посчитан сервером."
    )


def slide_limits(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, PAPER)
    title(s, "Работа важнее игры", 10)
    text(s, M, BODY_Y_SHORT, 6.6, 0.7, "Сервер сам ограничивает, когда и сколько можно играть. Все значения меняет администратор.",
         size=15, color=MUTED)
    rows(s, M, 2.55, 6.6, 2.2, [
        ("Набор заданий", "половина всех заданий в день"),
        ("Пауза", "30 минут между заданиями"),
        ("Часы заданий", "с 08:00 до 20:00"),
        ("Без перерыва", "не больше 1 часа подряд"),
        ("Отдых", "30 минут, иначе игра закрыта до завтра"),
    ], row_h=0.66)
    text(s, M, 6.1, 6.6, 0.8,
         "После часа игры есть 5 минут, чтобы выйти. Мини-игру лимит не прерывает: отдых начнётся сразу "
         "после неё.", size=13, color=MUTED)
    h = 5.2
    w = mock.body_width(h)
    px = RIGHT - w
    mock.phone(s, "36-play-rest", px, BODY_Y_SHORT, h)
    mock.phone(s, "35-play-warning", px - 0.333 - w, BODY_Y_SHORT, h)
    s.notes_slide.notes_text_frame.text = (
        "Игра мотивирует приходить, но не отнимает рабочее время. Развлекательные уведомления — только "
        "вне рабочего времени."
    )


def slide_admin(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, MIST)
    title(s, "Доступ выдаёт\nадминистратор", 11)
    prose(s, M, BODY_Y, 4.0, [
        ("Регистрации нет. Аккаунт создаёт администратор, у каждого игрока есть офис и департамент."),
        ("Руководитель видит статистику своих департаментов и заводит сотрудников."),
        ("Офисы и их сети, командировки, лимиты, журнал подозрительных действий."),
    ])
    mock.browser(s, "admin-04-stats", 5.35, BODY_Y, 8.6)
    s.notes_slide.notes_text_frame.text = (
        "Админка — отдельное веб-приложение на React. Панели администратора и руководителя "
        "используют разные RPC."
    )


def slide_architecture(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, PAPER)
    title(s, "Как это устроено", 12)
    bw, bh, gap = 3.2, 1.3, 1.2
    col1, col2, col3 = M, M + bw + gap, M + 2 * (bw + gap)
    row1, row2 = BODY_Y_SHORT, 3.95
    mid1, mid2 = row1 + bh / 2, row2 + bh / 2
    node(s, col1, row1, bw, bh, "Телефон сотрудника", "Safari · Godot 4 · WebGL 2")
    node(s, col1, row2, bw, bh, "Экран у ресепшна", "Ротируемые коды «Вход» и «Выход»")
    node(s, col2, row1, bw, row2 + bh - row1, "Nakama + Go-модуль",
         "Задания и проверка присутствия\nКошелёк с леджером\nЛимиты рабочего времени\nАнтифрод и журнал",
         accent=True)
    node(s, col3, row1, bw, bh, "PostgreSQL", "Аккаунты, леджер, хранилище")
    node(s, col3, row2, bw, bh, "Админка", "React · TypeScript · nakama-js")
    link(s, col1 + bw, mid1, col2, mid1, "WebSocket")
    link(s, col1 + bw, mid2, col2, mid2, "HTTP-ключ")
    link(s, col2 + bw, mid1, col3, mid1, "SQL")
    link(s, col3, mid2, col2 + bw, mid2, "HTTPS")
    text(s, M, 5.85, 7.6, 0.6,
         "Клиент ничего не начисляет сам: он отправляет действие через RPC, сервер проверяет его и меняет "
         "состояние.", size=15, spacing=1.22)
    hairline(s, M, 6.6, RIGHT, 6.6)
    text(s, M, 6.75, RIGHT - M, 0.3,
         "Тесты: GdUnit4, go test   ·   Линт: gdtoolkit, golangci-lint   ·   CI: GitHub Actions   ·   "
         "Запуск: Docker Compose", size=11, color=MUTED)
    s.notes_slide.notes_text_frame.text = (
        "Сервер — единственный источник истины. Офлайн-режима нет: игра всегда работает через сеть."
    )


def node(s, x, y, w, h, head, body, accent=False):
    fill, ink, sub = (RED, WHITE, WHITE) if accent else (MIST, INK, MUTED)
    box(s, x, y, w, h, fill, radius=0.08)
    lines = [[(head, {"bold": True, "size": 17, "color": ink})], [(body, {"size": 13, "color": sub})]]
    text(s, x + 0.25, y + 0.2, w - 0.5, h - 0.4, lines, spacing=1.18, space_after=4, anchor=MSO_ANCHOR.MIDDLE)


def link(s, x1, y1, x2, y2, label):
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = INK
    line.line.width = Pt(1.25)
    tail = line.line._get_or_add_ln()
    end = tail.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    tail.append(end)
    no_effects(line)
    if label:
        lx = min(x1, x2)
        text(s, lx + 0.1, min(y1, y2) - 0.3, abs(x2 - x1) - 0.2, 0.25, label, size=10, color=MUTED,
             align=PP_ALIGN.CENTER)


def slide_final(prs, mock):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, RED)
    text(s, M, TITLE_Y, 5.6, 0.3, "Альфа Офис", size=13, bold=True, color=WHITE)
    text(s, M, 1.85, 4.9, 3.2, "Игра работает только в Wi-Fi офиса. Поэтому в неё играют в офисе.",
         size=36, bold=True, color=WHITE, spacing=1.05)
    text(s, M, 6.62, 5.0, 0.3, "Студенческий концепт-проект, УрФУ, 2026", size=13, color=WHITE)
    h = 6.6
    w = mock.body_width(h)
    xb = RIGHT - w
    mock.phone(s, "44-office-colleagues", xb - 0.35 - w, 1.35, h)
    mock.phone(s, "11-go-home", xb, 2.05, h)
    s.notes_slide.notes_text_frame.text = "Спасибо. Готовы показать игру на телефоне."


def build(mock_dir: str) -> None:
    prs = Presentation()
    prs.slide_width = Emu(int(SLIDE_W * 914400))
    prs.slide_height = Emu(int(SLIDE_H * 914400))
    prs.core_properties.title = "Альфа Офис"
    prs.core_properties.author = "Альфа Офис"
    mock = Mockups(Path(mock_dir))
    for make in (slide_cover, slide_idea, slide_day, slide_tasks, slide_ar, slide_colleagues,
                 slide_presence, slide_economy, slide_shop, slide_limits, slide_admin,
                 slide_architecture, slide_final):
        make(prs, mock)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build(sys.argv[1])
